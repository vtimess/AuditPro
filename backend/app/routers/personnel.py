from datetime import date, datetime
from io import BytesIO
from urllib.parse import quote
from zoneinfo import ZoneInfo

from fastapi import APIRouter, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import ValidationError
from sqlalchemy import or_, select
from sqlalchemy.exc import IntegrityError

from app.core.id_card import (
    birth_date_from_id_card,
    decrypt_id_card,
    encrypt_id_card,
    hash_id_card,
    mask_id_card,
)
from app.dependencies import CurrentUser, DbSession
from app.models.business import Personnel
from app.schemas.common import ApiResponse
from app.schemas.personnel import PersonnelImportResult, PersonnelInput, PersonnelView, calculate_age, china_today


router = APIRouter(prefix="/personnel", tags=["人员管理"])

EXPORT_HEADERS = [
    "序号", "单位", "部门", "职务岗位", "在职状态", "姓名", "性别", "身份证号", "出生日期", "年龄",
    "民族", "学历", "基础病", "是否吸烟", "是否喝酒", "籍贯", "政治面貌", "入党时间", "是否参军",
    "参军工作时间", "进港时间", "联系电话", "家庭住址", "紧急联系人", "紧急联系人电话", "保险种类", "意外保险额度",
]
IMPORT_FIELDS = {
    "职务岗位": "position", "在职状态": "employment_status", "姓名": "person_name", "性别": "gender",
    "身份证号": "id_card", "民族": "nationality", "学历": "education", "基础病": "chronic_disease",
    "是否吸烟": "smokes", "是否喝酒": "drinks_alcohol", "籍贯": "native_place", "政治面貌": "political_status",
    "入党时间": "joined_party_at", "是否参军": "served_in_military", "参军工作时间": "military_service_time",
    "进港时间": "port_entry_date", "联系电话": "mobile", "家庭住址": "home_address",
    "紧急联系人": "emergency_contact", "紧急联系人电话": "emergency_mobile", "保险种类": "insurance_type",
}
REQUIRED_IMPORT_HEADERS = {"职务岗位", "在职状态", "姓名", "性别", "身份证号", "联系电话", "保险种类"}
BOOLEAN_FIELDS = {"smokes", "drinks_alcohol", "served_in_military"}
DATE_FIELDS = {"joined_party_at", "port_entry_date"}


def to_view(record: Personnel, reveal: bool = False) -> PersonnelView:
    id_card = decrypt_id_card(record.id_card_ciphertext)
    return PersonnelView(
        id=record.id, staff_no=record.staff_no or f"sy{record.id:03d}", unit_name=record.unit_name,
        department_name=record.department_name, position=record.position,
        employment_status=record.employment_status, person_name=record.person_name, gender=record.gender,
        id_card=id_card if reveal else mask_id_card(id_card), birth_date=record.birth_date,
        age=calculate_age(record.birth_date), nationality=record.nationality, education=record.education or "小学",
        chronic_disease=record.chronic_disease or "否",
        has_chronic_disease=record.has_chronic_disease, smokes=record.smokes,
        drinks_alcohol=record.drinks_alcohol, native_place=record.native_place,
        political_status=record.political_status, joined_party_at=record.joined_party_at,
        served_in_military=record.served_in_military, military_service_time=record.military_service_time,
        port_entry_date=record.port_entry_date, mobile=record.mobile, home_address=record.home_address,
        emergency_contact=record.emergency_contact, emergency_mobile=record.emergency_mobile,
        insurance_type=record.insurance_type, accident_insurance_limit=record.accident_insurance_limit,
        created_at=record.created_at,
    )


def build_record(payload: PersonnelInput, current_user: CurrentUser) -> Personnel:
    return Personnel(
        unit_name="泗阳队", department_name="泗阳劳务", position=payload.position,
        employment_status=payload.employment_status, person_name=payload.person_name.strip(), gender=payload.gender,
        id_card_ciphertext=encrypt_id_card(payload.id_card), id_card_hash=hash_id_card(payload.id_card),
        id_card_last4=payload.id_card[-4:], birth_date=birth_date_from_id_card(payload.id_card),
        nationality=payload.nationality.strip(), education=payload.education,
        chronic_disease=payload.chronic_disease,
        has_chronic_disease=payload.chronic_disease != "否", smokes=payload.smokes,
        drinks_alcohol=payload.drinks_alcohol, native_place=payload.native_place,
        political_status=payload.political_status,
        joined_party_at=None if payload.political_status == "群众" else payload.joined_party_at,
        served_in_military=payload.served_in_military,
        military_service_time=payload.military_service_time if payload.served_in_military else None,
        port_entry_date=payload.port_entry_date, mobile=payload.mobile, home_address=payload.home_address,
        emergency_contact=payload.emergency_contact, emergency_mobile=payload.emergency_mobile,
        insurance_type=payload.insurance_type, accident_insurance_limit="150+", created_by=current_user.id,
    )


def finalize_staff_no(db: DbSession, record: Personnel) -> None:
    db.flush()
    record.staff_no = f"sy{record.id:03d}"


@router.post("", response_model=ApiResponse[PersonnelView])
def create_personnel(payload: PersonnelInput, current_user: CurrentUser, db: DbSession):
    record = build_record(payload, current_user)
    db.add(record)
    try:
        finalize_staff_no(db, record)
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="该身份证号已存在于人员花名册") from exc
    db.refresh(record)
    return ApiResponse(message="人员信息保存成功", data=to_view(record))


@router.get("", response_model=ApiResponse[list[PersonnelView]])
def list_personnel(
    current_user: CurrentUser,
    db: DbSession,
    keyword: str = Query(default="", max_length=64),
    employment_status: str = Query(default="", max_length=20),
):
    query = select(Personnel).where(Personnel.record_status == "active")
    if keyword:
        query = query.where(or_(
            Personnel.person_name.contains(keyword), Personnel.mobile.contains(keyword),
            Personnel.id_card_last4.contains(keyword),
        ))
    if employment_status:
        if employment_status not in {"在职", "离职"}:
            raise HTTPException(status_code=400, detail="在职状态不正确")
        query = query.where(Personnel.employment_status == employment_status)
    records = db.scalars(query.order_by(Personnel.id.asc()).limit(2000)).all()
    return ApiResponse(data=[to_view(item) for item in records])


def workbook_response(workbook: Workbook, filename: str) -> StreamingResponse:
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


def style_header(sheet) -> None:
    fill = PatternFill("solid", fgColor="176C9B")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


@router.get("/template")
def download_template(current_user: CurrentUser):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "人员导入"
    sheet.append(EXPORT_HEADERS)
    style_header(sheet)
    widths = [12, 12, 14, 16, 12, 12, 10, 22, 14, 10, 12, 12, 12, 12, 12, 18, 16, 14, 12, 20, 14, 18, 32, 16, 20, 22, 16]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index) if index <= 26 else "AA"].width = width
    validations = {
        "D": '"经理,安全员,会计,带班长,一班班组长,二班班组长,装卸工"',
        "E": '"在职,离职"', "G": '"男,女"',
        "L": '"小学,初中,高中,中专,大专,本科,硕士及以上"',
        "M": '"否,高血压,糖尿病"', "N": '"否,是"',
        "O": '"否,是"', "S": '"否,是"', "Z": '"大港+工伤险,大港+雇主责任险"',
        "Q": '"群众,共青团员,中共党员,其他"',
    }
    for column, formula in validations.items():
        validation = DataValidation(type="list", formula1=formula, allow_blank=False)
        sheet.add_data_validation(validation)
        validation.add(f"{column}2:{column}2000")
    notes = workbook.create_sheet("填写说明")
    notes.append(["字段", "填写说明"])
    notes.append(["序号、出生日期、年龄", "无需填写，由系统自动生成"])
    notes.append(["单位、部门、意外保险额度", "固定为泗阳队、泗阳劳务、150+，导入时由系统统一设置"])
    notes.append(["身份证号", "必须是通过校验码验证的18位身份证号，单元格建议设为文本格式"])
    notes.append(["默认值", "学历默认为小学；基础病默认为否；政治面貌默认为群众；进港时间默认为导入当天"])
    notes.append(["入党时间", "政治面貌为群众时无需填写，系统统一按“无”处理"])
    notes.append(["联系电话", "必须填写以1开头的11位中国大陆手机号码"])
    notes.append(["必填字段", "职务岗位、在职状态、姓名、性别、身份证号、联系电话、保险种类"])
    style_header(notes)
    notes.column_dimensions["A"].width = 28
    notes.column_dimensions["B"].width = 88
    return workbook_response(workbook, "人员花名册导入模板.xlsx")


@router.get("/export")
def export_personnel(current_user: CurrentUser, db: DbSession):
    records = db.scalars(
        select(Personnel).where(Personnel.record_status == "active").order_by(Personnel.id.asc())
    ).all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "泗阳队花名册"
    sheet.append(EXPORT_HEADERS)
    for record in records:
        sheet.append([
            record.staff_no, record.unit_name, record.department_name, record.position, record.employment_status,
            record.person_name, record.gender, decrypt_id_card(record.id_card_ciphertext), record.birth_date.isoformat(),
            calculate_age(record.birth_date), record.nationality, record.education or "小学",
            record.chronic_disease or "否", "是" if record.smokes else "否",
            "是" if record.drinks_alcohol else "否", record.native_place or "", record.political_status or "",
            "无" if record.political_status == "群众" else (
                record.joined_party_at.isoformat() if record.joined_party_at else ""
            ),
            "是" if record.served_in_military else "否", record.military_service_time or "",
            record.port_entry_date.isoformat() if record.port_entry_date else "", record.mobile,
            record.home_address or "", record.emergency_contact or "", record.emergency_mobile or "",
            record.insurance_type, record.accident_insurance_limit,
        ])
    style_header(sheet)
    widths = [12, 12, 14, 16, 12, 12, 10, 22, 14, 10, 12, 12, 12, 12, 12, 18, 16, 14, 12, 20, 14, 18, 32, 16, 20, 22, 16]
    from openpyxl.utils import get_column_letter
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    timestamp = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y%m%d%H%M%S")
    return workbook_response(workbook, f"泗阳队花名册{timestamp}.xlsx")


def cell_text(value) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def cell_boolean(value) -> bool:
    text = cell_text(value)
    if text in {None, "否", "0", "否（默认）", "FALSE", "False", "false"}:
        return False
    if text in {"是", "1", "TRUE", "True", "true"}:
        return True
    raise ValueError("必须填写是或否")


def cell_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip().replace("/", "-").replace(".", "-")
    return datetime.strptime(text, "%Y-%m-%d").date()


@router.post("/import", response_model=ApiResponse[PersonnelImportResult])
async def import_personnel(current_user: CurrentUser, db: DbSession, file: UploadFile = File(...)):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 格式的Excel文件")
    content = await file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="Excel文件不能超过5MB")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook["人员导入"] if "人员导入" in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Excel文件无法读取或没有数据") from exc
    headers = [cell_text(value) or "" for value in header_row]
    missing_headers = REQUIRED_IMPORT_HEADERS - set(headers)
    if missing_headers:
        raise HTTPException(status_code=400, detail=f"Excel缺少必填列：{'、'.join(sorted(missing_headers))}")

    existing_hashes = set(db.scalars(select(Personnel.id_card_hash)).all())
    pending_hashes: set[str] = set()
    valid_payloads: list[PersonnelInput] = []
    errors: list[str] = []
    for row_number, row in enumerate(rows, 2):
        if not any(value not in (None, "") for value in row):
            continue
        raw = {headers[index]: value for index, value in enumerate(row) if index < len(headers)}
        data = {}
        try:
            for header, field in IMPORT_FIELDS.items():
                value = raw.get(header)
                if field in BOOLEAN_FIELDS:
                    data[field] = cell_boolean(value)
                elif field in DATE_FIELDS:
                    data[field] = cell_date(value)
                else:
                    data[field] = cell_text(value)
            data["employment_status"] = data.get("employment_status") or "在职"
            data["nationality"] = data.get("nationality") or "汉族"
            data["education"] = data.get("education") or "小学"
            data["chronic_disease"] = data.get("chronic_disease") or "否"
            data["political_status"] = data.get("political_status") or "群众"
            data["port_entry_date"] = data.get("port_entry_date") or china_today()
            payload = PersonnelInput.model_validate(data)
            id_hash = hash_id_card(payload.id_card)
            if id_hash in existing_hashes or id_hash in pending_hashes:
                raise ValueError("身份证号已存在于人员花名册")
            pending_hashes.add(id_hash)
            valid_payloads.append(payload)
        except (ValueError, ValidationError) as exc:
            if isinstance(exc, ValidationError):
                message = str(exc.errors()[0].get("msg", "数据格式不正确")).replace("Value error, ", "")
            else:
                message = str(exc)
            errors.append(f"第{row_number}行：{message}")

    try:
        for payload in valid_payloads:
            record = build_record(payload, current_user)
            db.add(record)
            finalize_staff_no(db, record)
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="导入数据存在重复身份证号，请检查后重试") from exc
    result = PersonnelImportResult(
        success_count=len(valid_payloads), failure_count=len(errors), errors=errors[:50],
    )
    return ApiResponse(message=f"成功导入{result.success_count}人，失败{result.failure_count}人", data=result)


@router.get("/{personnel_id}", response_model=ApiResponse[PersonnelView])
def get_personnel(personnel_id: int, current_user: CurrentUser, db: DbSession):
    record = db.get(Personnel, personnel_id)
    if not record or record.record_status != "active":
        raise HTTPException(status_code=404, detail="人员不存在")
    return ApiResponse(data=to_view(record, reveal=True))


@router.put("/{personnel_id}", response_model=ApiResponse[PersonnelView])
def update_personnel(personnel_id: int, payload: PersonnelInput, current_user: CurrentUser, db: DbSession):
    record = db.get(Personnel, personnel_id)
    if not record or record.record_status != "active":
        raise HTTPException(status_code=404, detail="人员不存在")
    values = payload.model_dump(exclude={"id_card"})
    for key, value in values.items():
        setattr(record, key, value)
    record.unit_name = "泗阳队"
    record.department_name = "泗阳劳务"
    record.accident_insurance_limit = "150+"
    record.has_chronic_disease = payload.chronic_disease != "否"
    record.joined_party_at = None if payload.political_status == "群众" else payload.joined_party_at
    record.military_service_time = payload.military_service_time if payload.served_in_military else None
    record.id_card_ciphertext = encrypt_id_card(payload.id_card)
    record.id_card_hash = hash_id_card(payload.id_card)
    record.id_card_last4 = payload.id_card[-4:]
    record.birth_date = birth_date_from_id_card(payload.id_card)
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="该身份证号已存在于人员花名册") from exc
    db.refresh(record)
    return ApiResponse(message="人员信息修改成功", data=to_view(record))


@router.delete("/{personnel_id}", response_model=ApiResponse[bool])
def delete_personnel(personnel_id: int, current_user: CurrentUser, db: DbSession):
    record = db.get(Personnel, personnel_id)
    if not record or record.record_status != "active":
        raise HTTPException(status_code=404, detail="人员不存在")
    record.record_status = "deleted"
    db.commit()
    return ApiResponse(message="人员已从花名册删除", data=True)
