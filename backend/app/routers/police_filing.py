from datetime import date
from io import BytesIO
from urllib.parse import quote

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError

from app.core.id_card import age_from_birth_date, birth_date_from_id_card, decrypt_id_card, encrypt_id_card, hash_id_card, mask_id_card
from app.dependencies import CurrentUser, DbSession
from app.models.business import PoliceRegistration
from app.schemas.business import FIXED_FILING_COMPANY, PoliceRegistrationInput, PoliceRegistrationView
from app.schemas.common import ApiResponse


router = APIRouter(prefix="/police-filings", tags=["安全备案"])
MANAGER_ROLES = {"group_leader", "regional_manager", "regional_safety_manager", "project_leader"}


def scoped_query(current_user: CurrentUser):
    query = select(PoliceRegistration).where(PoliceRegistration.record_status == "active")
    if current_user.role_code not in MANAGER_ROLES:
        query = query.where(PoliceRegistration.created_by == current_user.id)
    return query


def to_view(record: PoliceRegistration, reveal: bool = False) -> PoliceRegistrationView:
    id_card = decrypt_id_card(record.id_card_ciphertext)
    return PoliceRegistrationView(
        id=record.id,
        person_name=record.person_name,
        gender=record.gender,
        id_card=id_card if reveal else mask_id_card(id_card),
        age=age_from_birth_date(birth_date_from_id_card(id_card)),
        company_name=FIXED_FILING_COMPANY,
        region=record.region,
        mobile=record.mobile,
        filing_status=record.filing_status,
        filing_date=record.filing_date,
        remarks=record.remarks,
        created_at=record.created_at,
    )


@router.post("", response_model=ApiResponse[PoliceRegistrationView])
def create_filing(payload: PoliceRegistrationInput, current_user: CurrentUser, db: DbSession):
    record = PoliceRegistration(
        created_by=current_user.id,
        person_name=payload.person_name.strip(),
        gender=payload.gender,
        id_card_ciphertext=encrypt_id_card(payload.id_card),
        id_card_hash=hash_id_card(payload.id_card),
        id_card_last4=payload.id_card[-4:],
        company_name=FIXED_FILING_COMPANY,
        region=payload.region.strip(),
        mobile=payload.mobile,
        filing_status=payload.filing_status,
        filing_date=payload.filing_date,
        remarks=payload.remarks,
    )
    db.add(record)
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="该身份证号已登记") from exc
    db.refresh(record)
    return ApiResponse(message="备案人员保存成功", data=to_view(record))


@router.get("", response_model=ApiResponse[list[PoliceRegistrationView]])
def list_filings(
    current_user: CurrentUser,
    db: DbSession,
    keyword: str = Query(default="", max_length=64),
    region: str = Query(default="", max_length=64),
    filing_status: str = Query(default="", max_length=20),
):
    query = scoped_query(current_user)
    if keyword:
        query = query.where(
            PoliceRegistration.person_name.contains(keyword)
            | PoliceRegistration.company_name.contains(keyword)
            | PoliceRegistration.id_card_last4.contains(keyword)
        )
    if region:
        query = query.where(PoliceRegistration.region == region)
    if filing_status:
        query = query.where(PoliceRegistration.filing_status == filing_status)
    records = db.scalars(query.order_by(PoliceRegistration.created_at.desc()).limit(500)).all()
    return ApiResponse(data=[to_view(item) for item in records])


@router.get("/export")
def export_filings(
    current_user: CurrentUser,
    db: DbSession,
    keyword: str = Query(default="", max_length=64),
    region: str = Query(default="", max_length=64),
    filing_status: str = Query(default="", max_length=20),
):
    query = scoped_query(current_user)
    if keyword:
        query = query.where(
            PoliceRegistration.person_name.contains(keyword)
            | PoliceRegistration.company_name.contains(keyword)
            | PoliceRegistration.id_card_last4.contains(keyword)
        )
    if region:
        query = query.where(PoliceRegistration.region == region)
    if filing_status:
        query = query.where(PoliceRegistration.filing_status == filing_status)
    records = db.scalars(query.order_by(PoliceRegistration.created_at.desc())).all()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "拟录用职工名单"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:H1")
    sheet["A1"] = "泗阳县星光装卸有限公司拟录用职工名单"
    sheet["A1"].font = Font(name="宋体", size=20)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 42

    headers = ["序号", "姓名", "性别", "年龄", "居民身份证号", "联系方式", "用工公司", "备注"]
    for column, value in enumerate(headers, 1):
        cell = sheet.cell(row=2, column=column, value=value)
        cell.font = Font(name="宋体", size=12)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[2].height = 30

    for index, record in enumerate(records, 1):
        id_card = decrypt_id_card(record.id_card_ciphertext)
        row = index + 2
        values = [
            index, record.person_name, record.gender,
            age_from_birth_date(birth_date_from_id_card(id_card)), id_card,
            record.mobile or "", FIXED_FILING_COMPANY, record.remarks or "",
        ]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.font = Font(name="宋体", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[row].height = 26

    footer_row = max(12, len(records) + 3)
    for row in range(3, footer_row):
        sheet.row_dimensions[row].height = 26
    thin = Side(style="thin", color="000000")
    table_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in sheet.iter_rows(min_row=1, max_row=footer_row - 1, min_col=1, max_col=8):
        for cell in row:
            cell.border = table_border

    sheet.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=4)
    sheet.merge_cells(start_row=footer_row, start_column=5, end_row=footer_row, end_column=8)
    sheet.cell(row=footer_row, column=1, value="派出所备案情况：")
    today = date.today()
    sheet.cell(row=footer_row, column=5, value=f"日期{today.year}年{today.month}月{today.day}日")
    sheet.cell(row=footer_row, column=1).alignment = Alignment(horizontal="left", vertical="center")
    sheet.cell(row=footer_row, column=5).alignment = Alignment(horizontal="center", vertical="center")
    sheet.cell(row=footer_row, column=1).font = Font(name="宋体", size=11)
    sheet.cell(row=footer_row, column=5).font = Font(name="宋体", size=11)
    sheet.row_dimensions[footer_row].height = 28

    widths = (8, 13, 8, 8, 28, 18, 25, 20)
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_area = f"A1:H{footer_row}"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = quote(f"拟录用职工名单_{date.today().isoformat()}.xlsx")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


@router.get("/{filing_id}", response_model=ApiResponse[PoliceRegistrationView])
def get_filing(filing_id: int, current_user: CurrentUser, db: DbSession):
    record = db.get(PoliceRegistration, filing_id)
    if not record or record.record_status != "active":
        raise HTTPException(status_code=404, detail="备案人员不存在")
    if current_user.role_code not in MANAGER_ROLES and record.created_by != current_user.id:
        raise HTTPException(status_code=403, detail="无权查看该记录")
    return ApiResponse(data=to_view(record, reveal=True))


@router.put("/{filing_id}", response_model=ApiResponse[PoliceRegistrationView])
def update_filing(filing_id: int, payload: PoliceRegistrationInput, current_user: CurrentUser, db: DbSession):
    record = db.get(PoliceRegistration, filing_id)
    if not record or record.record_status != "active":
        raise HTTPException(status_code=404, detail="备案人员不存在")
    if current_user.role_code not in MANAGER_ROLES and record.created_by != current_user.id:
        raise HTTPException(status_code=403, detail="无权修改该记录")
    for key in ("person_name", "gender", "region", "mobile", "filing_status", "filing_date", "remarks"):
        setattr(record, key, getattr(payload, key))
    record.company_name = FIXED_FILING_COMPANY
    record.id_card_ciphertext = encrypt_id_card(payload.id_card)
    record.id_card_hash = hash_id_card(payload.id_card)
    record.id_card_last4 = payload.id_card[-4:]
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="该身份证号已登记") from exc
    db.refresh(record)
    return ApiResponse(message="备案信息修改成功", data=to_view(record))


@router.delete("/{filing_id}", response_model=ApiResponse[bool])
def delete_filing(filing_id: int, current_user: CurrentUser, db: DbSession):
    record = db.get(PoliceRegistration, filing_id)
    if not record or record.record_status != "active":
        raise HTTPException(status_code=404, detail="备案人员不存在")
    if current_user.role_code not in MANAGER_ROLES and record.created_by != current_user.id:
        raise HTTPException(status_code=403, detail="无权删除该记录")
    record.record_status = "deleted"
    db.commit()
    return ApiResponse(message="备案信息已删除", data=True)
