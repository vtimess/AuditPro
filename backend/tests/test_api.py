from __future__ import annotations

import os
import re
from datetime import date, datetime, timedelta
from io import BytesIO
from types import SimpleNamespace
from urllib.parse import unquote
from zoneinfo import ZoneInfo
from zipfile import ZIP_DEFLATED, ZipFile

os.environ.setdefault("APP_ENV", "development")
os.environ.setdefault("AUTO_CREATE_TABLES", "false")
os.environ.setdefault("DB_HOST", "127.0.0.1")
os.environ.setdefault("DB_USER", "test")
os.environ.setdefault("DB_PASSWORD", "test")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402
from sqlalchemy.pool import StaticPool  # noqa: E402

from app.database import Base, get_db  # noqa: E402
from app.main import app  # noqa: E402
from app.routers import auth as auth_router  # noqa: E402
from app.routers import emergency_drill as emergency_drill_router  # noqa: E402


engine = create_engine(
    "sqlite://",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSession = sessionmaker(bind=engine, expire_on_commit=False)
Base.metadata.create_all(engine)


def override_db():
    db = TestingSession()
    try:
        yield db
    finally:
        db.close()


app.dependency_overrides[get_db] = override_db
client = TestClient(app)


def login() -> tuple[str, dict]:
    response = client.post(
        "/api/v1/auth/wechat-login",
        json={"code": "test-code", "nickname": "测试用户", "avatar_url": ""},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["code"] == 0
    return body["data"]["access_token"], body["data"]["user"]


def test_health():
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_login_home_workbench_and_profile(monkeypatch):
    monkeypatch.setattr(
        auth_router,
        "exchange_code",
        lambda _: {"openid": "test_openid", "unionid": None},
    )
    token, user = login()
    headers = {"Authorization": f"Bearer {token}"}
    assert user["role_code"] == "ordinary_user"

    home = client.get("/api/v1/home", headers=headers)
    assert home.status_code == 200
    assert home.json()["data"]["user"]["real_name"] == "测试用户"
    assert len(home.json()["data"]["metrics"]) == 4

    workbench = client.get("/api/v1/workbench", headers=headers)
    assert workbench.status_code == 200
    assert len(workbench.json()["data"]["modules"]) >= 5
    assert any(item["key"] == "personnel_management" and item["enabled"] for item in workbench.json()["data"]["modules"])
    assert any(item["key"] == "emergency_drill" and item["enabled"] for item in workbench.json()["data"]["modules"])
    talk_module = next(item for item in workbench.json()["data"]["modules"] if item["key"] == "heart_to_heart_talk")
    assert talk_module["enabled"] is False
    assert talk_module["status_text"] == "规划中"

    profile = client.get("/api/v1/profile", headers=headers)
    assert profile.status_code == 200

    updated = client.patch(
        "/api/v1/profile",
        headers=headers,
        json={"real_name": "张三", "mobile": "13800000000"},
    )
    assert updated.status_code == 200
    assert updated.json()["data"]["real_name"] == "张三"
    assert updated.json()["data"]["mobile"] == "13800000000"


def test_protected_endpoint_requires_token():
    response = client.get("/api/v1/profile")
    assert response.status_code == 401


def test_attendance_punch_calendar_and_supplement(monkeypatch):
    monkeypatch.setattr(
        auth_router,
        "exchange_code",
        lambda _: {"openid": "attendance_openid", "unionid": None},
    )
    token, _ = login()
    headers = {"Authorization": f"Bearer {token}"}

    today = client.get("/api/v1/attendance/today", headers=headers)
    assert today.status_code == 200
    assert today.json()["data"]["current_action"] == "check_in"

    payload = {"latitude": 31.2304, "longitude": 121.4737, "accuracy": 10, "location_text": "测试位置"}
    check_in = client.post("/api/v1/attendance/punch", headers=headers, json=payload)
    assert check_in.status_code == 200
    assert check_in.json()["data"]["punch_type"] == "check_in"

    check_out = client.post("/api/v1/attendance/punch", headers=headers, json=payload)
    assert check_out.status_code == 200
    assert check_out.json()["data"]["punch_type"] == "check_out"

    month = date.today().strftime("%Y-%m")
    calendar_response = client.get(f"/api/v1/attendance/month?month={month}", headers=headers)
    assert calendar_response.status_code == 200
    today_item = next(item for item in calendar_response.json()["data"]["days"] if item["date"] == date.today().isoformat())
    assert today_item["status"] == "complete"

    supplement_day = date.today() - timedelta(days=1)
    supplement = client.post(
        "/api/v1/attendance/supplements",
        headers=headers,
        json={
            "work_date": supplement_day.isoformat(),
            "punch_type": "check_in",
            "punch_time": "09:00:00",
            "reason": "外出作业忘记打卡",
        },
    )
    assert supplement.status_code == 200
    assert supplement.json()["data"]["review_status"] == "pending"


def test_police_purchase_and_safety_modules(monkeypatch):
    monkeypatch.setattr(
        auth_router,
        "exchange_code",
        lambda _: {"openid": "business_openid", "unionid": None},
    )
    token, _ = login()
    headers = {"Authorization": f"Bearer {token}"}

    filing = client.post(
        "/api/v1/police-filings",
        headers=headers,
        json={
            "person_name": "李四", "gender": "男", "id_card": "11010519491231002X",
            "company_name": "泗阳星光装卸有限公司", "region": "镇江", "mobile": "13800000000",
            "filing_status": "pending",
        },
    )
    assert filing.status_code == 200, filing.text
    assert "*" in filing.json()["data"]["id_card"]
    assert filing.json()["data"]["age"] > 0
    export = client.get("/api/v1/police-filings/export", headers=headers)
    assert export.status_code == 200
    assert export.headers["content-type"].startswith("application/vnd.openxmlformats")
    export_book = load_workbook(BytesIO(export.content))
    export_sheet = export_book.active
    assert export_sheet["A1"].value == "泗阳县星光装卸有限公司拟录用职工名单"
    assert [export_sheet.cell(row=2, column=i).value for i in range(1, 9)] == [
        "序号", "姓名", "性别", "年龄", "居民身份证号", "联系方式", "用工公司", "备注",
    ]
    assert export_sheet["G3"].value == "泗阳星光装卸有限公司"
    assert export_sheet["A12"].value == "派出所备案情况："
    assert export_sheet["E12"].value == f"日期{date.today().year}年{date.today().month}月{date.today().day}日"

    rejected_purchase = client.post(
        "/api/v1/consumable-purchases",
        headers=headers,
        json={
            "category": "劳保用品", "item_name": "安全帽", "quantity": "10", "unit": "个",
            "unit_price": "20.00", "total_amount": "180.00", "purchase_date": date.today().isoformat(),
            "region": "镇江",
        },
    )
    assert rejected_purchase.status_code == 422
    purchase = client.post(
        "/api/v1/consumable-purchases",
        headers=headers,
        json={
            "category": "劳保用品", "item_name": "安全帽", "quantity": "10", "unit": "个",
            "unit_price": "20.00", "total_amount": "180.00", "total_adjustment_reason": "供应商优惠",
            "purchase_date": date.today().isoformat(), "region": "镇江",
        },
    )
    assert purchase.status_code == 200, purchase.text
    assert purchase.json()["data"]["is_manual_total"] is True
    assert purchase.json()["data"]["calculated_total"] == "200.00"
    purchase_id = purchase.json()["data"]["id"]
    selected_export = client.get(f"/api/v1/consumable-purchases/export?ids={purchase_id}", headers=headers)
    assert selected_export.status_code == 200
    assert selected_export.headers["content-type"].startswith("application/vnd.openxmlformats")

    inspection = client.post(
        "/api/v1/safety-inspections",
        headers=headers,
        json={
            "inspection_project": "一号码头", "region": "镇江", "inspection_area": "装卸区",
            "inspected_at": "2026-08-13T09:30:00", "inspectors": "测试用户",
            "site_readiness": "rectification", "hazards": [{
                "hazard_name": "临边防护缺失", "category": "作业环境", "risk_level": "major",
                "description": "临边区域未设置防护栏", "responsible_person": "测试用户",
                "rectification_deadline": "2026-08-20", "rectification_requirement": "设置防护栏和警示标志",
                "immediate_stop": True,
            }],
        },
    )
    assert inspection.status_code == 200, inspection.text
    hazard = inspection.json()["data"]["hazards"][0]
    assert hazard["risk_level"] == "major"
    rectified = client.post(
        f"/api/v1/safety-inspections/hazards/{hazard['id']}/rectify",
        headers=headers,
        json={"description": "已安装防护栏和警示标志"},
    )
    assert rectified.status_code == 200
    assert rectified.json()["data"]["status"] == "pending_review"


def test_personnel_crud_import_and_export(monkeypatch):
    monkeypatch.setattr(
        auth_router,
        "exchange_code",
        lambda _: {"openid": "personnel_openid", "unionid": None},
    )
    token, _ = login()
    headers = {"Authorization": f"Bearer {token}"}
    payload = {
        "position": "装卸工", "employment_status": "在职", "person_name": "王五", "gender": "男",
        "id_card": "11010519491231002X", "nationality": "汉族",
        "smokes": False, "drinks_alcohol": False,
        "native_place": "江苏泗阳", "served_in_military": False,
        "mobile": "13900000000", "home_address": "泗阳县测试路1号",
        "emergency_contact": "王六", "emergency_mobile": "13700000000", "insurance_type": "大港+工伤险",
    }
    created = client.post("/api/v1/personnel", headers=headers, json=payload)
    assert created.status_code == 200, created.text
    person = created.json()["data"]
    assert person["staff_no"] == "sy001"
    assert person["unit_name"] == "泗阳队"
    assert person["department_name"] == "泗阳劳务"
    assert person["birth_date"] == "1949-12-31"
    assert person["education"] == "小学"
    assert person["chronic_disease"] == "否"
    assert person["political_status"] == "群众"
    assert person["joined_party_at"] is None
    china_date = datetime.now(ZoneInfo("Asia/Shanghai")).date()
    assert person["port_entry_date"] == china_date.isoformat()
    assert person["accident_insurance_limit"] == "150+"

    invalid_mobile = client.post(
        "/api/v1/personnel", headers=headers, json={**payload, "mobile": "1390000000"},
    )
    assert invalid_mobile.status_code == 422

    roster = client.get("/api/v1/personnel?keyword=王五&employment_status=在职", headers=headers)
    assert roster.status_code == 200
    assert roster.json()["data"][0]["id_card"].count("*") > 0

    payload["employment_status"] = "离职"
    updated = client.put(f"/api/v1/personnel/{person['id']}", headers=headers, json=payload)
    assert updated.status_code == 200
    assert updated.json()["data"]["employment_status"] == "离职"

    template = client.get("/api/v1/personnel/template", headers=headers)
    assert template.status_code == 200
    assert template.headers["content-type"].startswith("application/vnd.openxmlformats")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "人员导入"
    sheet.append(["职务岗位", "在职状态", "姓名", "性别", "身份证号", "联系电话", "保险种类"])
    sheet.append(["安全员", "在职", "赵六", "女", "110105198806150032", "13600000000", "大港+雇主责任险"])
    output = BytesIO()
    workbook.save(output)
    imported = client.post(
        "/api/v1/personnel/import", headers=headers,
        files={"file": ("人员导入.xlsx", output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["data"]["success_count"] == 1
    assert imported.json()["data"]["failure_count"] == 0

    exported = client.get("/api/v1/personnel/export", headers=headers)
    assert exported.status_code == 200
    assert exported.headers["content-type"].startswith("application/vnd.openxmlformats")
    disposition = unquote(exported.headers["content-disposition"])
    assert re.search(r"泗阳队花名册\d{14}\.xlsx", disposition)
    export_book = load_workbook(BytesIO(exported.content), data_only=True)
    export_sheet = export_book["泗阳队花名册"]
    assert export_sheet["L2"].value == "小学"
    assert export_sheet["M2"].value == "否"
    assert export_sheet["Q2"].value == "群众"
    assert export_sheet["R2"].value == "无"
    assert export_sheet["U2"].value == china_date.isoformat()

    deleted = client.delete(f"/api/v1/personnel/{person['id']}", headers=headers)
    assert deleted.status_code == 200


def test_emergency_drill_word_pdf_import_preview_and_export(monkeypatch, tmp_path):
    monkeypatch.setattr(
        auth_router,
        "exchange_code",
        lambda _: {"openid": "emergency_drill_openid", "unionid": None},
    )
    monkeypatch.setattr(emergency_drill_router, "settings", SimpleNamespace(file_root=str(tmp_path)))
    token, _ = login()
    headers = {"Authorization": f"Bearer {token}"}

    rejected = client.post(
        "/api/v1/emergency-drills/import", headers=headers,
        data={"original_name": "错误格式.xlsx"},
        files={"file": ("错误格式.xlsx", b"not-an-excel-file", "application/octet-stream")},
    )
    assert rejected.status_code == 400

    pdf_content = b"%PDF-1.4\n1 0 obj<</Type/Catalog>>endobj\n%%EOF"
    imported_pdf = client.post(
        "/api/v1/emergency-drills/import", headers=headers,
        data={"original_name": "消防应急演练方案.pdf"},
        files={"file": ("upload.pdf", pdf_content, "application/pdf")},
    )
    assert imported_pdf.status_code == 200, imported_pdf.text
    pdf_record = imported_pdf.json()["data"]
    assert pdf_record["file_name"] == "消防应急演练方案.pdf"
    assert pdf_record["file_type"] == "PDF"

    docx_output = BytesIO()
    with ZipFile(docx_output, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types></Types>")
        archive.writestr("word/document.xml", "<w:document></w:document>")
    docx_content = docx_output.getvalue()
    imported_docx = client.post(
        "/api/v1/emergency-drills/import", headers=headers,
        data={"original_name": "防汛防台演练方案.docx"},
        files={"file": ("upload.docx", docx_content, "application/octet-stream")},
    )
    assert imported_docx.status_code == 200, imported_docx.text
    docx_record = imported_docx.json()["data"]
    assert docx_record["file_type"] == "DOCX"

    doc_content = bytes.fromhex("D0CF11E0A1B11AE1") + b"legacy-word-document"
    imported_doc = client.post(
        "/api/v1/emergency-drills/import", headers=headers,
        data={"original_name": "设备事故演练方案.doc"},
        files={"file": ("upload.doc", doc_content, "application/msword")},
    )
    assert imported_doc.status_code == 200, imported_doc.text
    assert imported_doc.json()["data"]["file_type"] == "DOC"

    listed = client.get("/api/v1/emergency-drills?keyword=消防", headers=headers)
    assert listed.status_code == 200
    assert [item["file_name"] for item in listed.json()["data"]] == ["消防应急演练方案.pdf"]

    preview = client.get(pdf_record["preview_url"], headers=headers)
    assert preview.status_code == 200
    assert preview.content == pdf_content
    assert preview.headers["content-type"].startswith("application/pdf")

    exported = client.get(docx_record["export_url"], headers=headers)
    assert exported.status_code == 200
    assert exported.content == docx_content
    assert "防汛防台演练方案.docx" in unquote(exported.headers["content-disposition"])
